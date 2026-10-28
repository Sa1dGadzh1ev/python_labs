import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Проверяет CSV, считает ширины и построчно записывает XLSX."""
    csv_path = Path(csv_path)
    xlsx_path = Path(xlsx_path)

    if csv_path.suffix.lower() != ".csv":
        raise ValueError("Ожидался файл с расширением .csv")

    if xlsx_path.suffix.lower() != ".xlsx":
        raise ValueError("Ожидался файл с расширением .xlsx")

    with csv_path.open(encoding="utf-8", newline="") as f:
        reader = csv.reader(f, delimiter=",", strict=True)

        try:
            header = next(reader, None)

            if not header:
                raise ValueError("Пустой CSV или отсутствует заголовок")

            names = []
            widths = []

            for name in header:
                if not name.strip():
                    raise ValueError("Название колонки не должно быть пустым")

                if name in names:
                    raise ValueError("Названия колонок не должны повторяться")

                names.append(name)

                width = 8

                for line in name.split("\n"):
                    if len(line) + 2 > width:
                        width = len(line) + 2

                widths.append(width)

            for row in reader:
                if len(row) != len(header):
                    raise ValueError("Количество полей не совпадает с заголовком")

                for i in range(len(row)):
                    for line in row[i].split("\n"):
                        if len(line) + 2 > widths[i]:
                            widths[i] = len(line) + 2

        except csv.Error as error:
            raise ValueError("Некорректный CSV") from error

        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Sheet1")

        for i in range(len(widths)):
            letter = get_column_letter(i + 1)
            ws.column_dimensions[letter].width = widths[i]

        f.seek(0)
        reader = csv.reader(f, delimiter=",", strict=True)

        for row in reader:
            cells = []

            for value in row:
                cell = WriteOnlyCell(ws, value=value)
                cell.data_type = "s"
                cells.append(cell)

            ws.append(cells)

        xlsx_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(xlsx_path)
        wb.close()